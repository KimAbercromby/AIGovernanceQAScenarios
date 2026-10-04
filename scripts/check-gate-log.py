#!/usr/bin/env python3
"""Reproduce the Gate Log (AIG-DEC-04) workbook checks for scenarios G1 to G9 and H2 (F4b to F4d in the Register and I7 to I10 in the Agent Record, if they are in the same folder).

Usage: python3 scripts/check-gate-log.py <AIG-DEC-04_Gate_Log_Proposed.xlsx> [path/to/recalc.py]

Writes sample rows into a COPY of the workbook, recalculates it with LibreOffice
(through recalc.py, which opens the copy headless and saves the computed values),
reads each row-check cell back and writes results/workbook-results.json and
results/workbook-results.js. The original workbook is never changed.
"""
import datetime, json, os, shutil, subprocess, sys, tempfile
import openpyxl

src = sys.argv[1]
recalc = sys.argv[2] if len(sys.argv) > 2 else "/mnt/skills/public/xlsx/scripts/recalc.py"
tmp = os.path.join(tempfile.mkdtemp(), "gate-log-test.xlsx")
shutil.copy(src, tmp)
wb = openpyxl.load_workbook(tmp)
P, E, C = wb["Gate plan"], wb["Gate events"], wb["Conditions"]
g6, g1, g0 = "Gate 6 Deployment / go-live", "Gate 1 Strategic prioritisation", "Gate 0 Intake / opportunity"
def put(ws, r, **v):
    for k, x in v.items(): ws[f"{k}{r}"] = x
d = datetime.datetime(2026, 10, 1)
plan = {  # row: (scenario, values)
    4: ("G7", dict(A="PL-1", B="AIR-T001", C=g6, E="Required", F="triage", I="Planned", K="UC-1", L="UC-ID specific")),
    5: ("G1b", dict(A="PL-2", B="AIR-T001", C=g1, F="triage", I="Planned", K="UC-1", L="UC-ID specific")),
    6: ("G1c", dict(A="PL-3", B="AIR-T001", C=g1, E="Required", F="x", K="UC-2", L="UC-ID specific")),
    7: ("G1", dict(A="PL-4", B="AIR-T001", C="Gate 4 Procurement", E="Not applicable", F="x", I="Planned", K="UC-3", L="UC-ID specific")),
}
base = dict(B="AIR-T001", F=d, H="Officer", I="DEC-1", P="UC-ID specific")
events = {
    4: ("G7", dict(A="EV-1", C="PL-1", D=g6, E="Decision", G="Progress", O="UC-1", X="Yes")),
    5: ("G2", dict(A="EV-2", C="PL-1", D=g1, E="Decision", G="Progress", O="UC-1")),
    6: ("G2b", dict(A="EV-3", C="PL-1", D=g6, E="Decision", G="Progress", O="UC-9")),
    7: ("G3b", dict(A="EV-4", D=g6, E="Assurance opinion", G="Opinion only", O="UC-1")),
    8: ("G3", dict(A="EV-5", D=g0, E="Intake / registration", G="Opinion only", O="UC-1")),
    9: ("G4", dict(A="EV-6", D=g6, E="Decision", G="Progress", O="UC-1 UC-2")),
    10: ("G5", dict(A="EV-7", D=g6, E="Precautionary pause (containment)", G="Paused — pending decision", V="INC-1", W=datetime.datetime(2026, 9, 1), O="UC-1")),
    11: ("G7", dict(A="EV-8", D=g0, E="Intake / registration", G="No decision", O="UC-1")),
    12: ("G7", dict(A="EV-9", D=g6, E="Decision", G="Progress with condition", O="UC-1", X="Yes")),
    13: ("G9", dict(A="EV-10", D=g6, E="Decision", G="Progress", O="UC-1")),
    14: ("G9b", dict(A="EV-11", D=g6, E="Decision", G="Re-authorise", O="UC-1", X="No")),
    15: ("H2", dict(A="EV-12", D=g6, E="Precautionary pause (containment)", G="Paused — pending decision", V="INC-2", W=datetime.datetime(2026, 10, 20), O="UC-5")),
    16: ("H2", dict(A="EV-13", D=g6, E="Decision", G="Resume", O="UC-5", X="Yes")),
    17: ("H2b", dict(A="EV-14", D=g6, E="Decision", G="Resume", O="UC-1", X="Yes")),
    18: ("H2c", dict(A="EV-15", D=g6, E="Decision", G="Resume", O="UC-6", X="Yes")),
    19: ("H2", dict(A="EV-16", D=g6, E="Precautionary pause (containment)", G="Paused — pending decision", V="INC-3", W=datetime.datetime(2026, 10, 20), O="UC-7")),
    20: ("H2d", dict(A="EV-17", D=g6, E="Decision", G="Resume", O="UC-7", X="No")),
    21: ("H2e", dict(A="EV-18", D=g6, E="Assurance opinion", G="Resume", J="AO-1", O="UC-5")),
}
conds = {
    4: ("G6", dict(A="C-1", B="EV-9", C="AIR-T001", D="do x", E="Owner", F=datetime.datetime(2027, 1, 1), G="Unknown", K="UC-1", L="UC-ID specific")),
    5: ("G7", dict(A="C-2", B="EV-9", C="AIR-T001", D="do y", E="Owner", F=datetime.datetime(2027, 1, 1), G="Open", K="UC-1", L="UC-ID specific")),
    6: ("H2", dict(A="C-3", B="EV-13", C="AIR-T001", D="extra check after the pause", E="Owner", F=datetime.datetime(2027, 1, 1), G="Open", K="UC-5", L="UC-ID specific")),
}
for r, (_, v) in plan.items(): put(P, r, **v)
for r, (_, v) in events.items(): put(E, r, **{**base, **v})
for r, (_, v) in conds.items(): put(C, r, **v)
wb.save(tmp)
subprocess.run([sys.executable, os.path.basename(recalc), tmp], cwd=os.path.dirname(recalc), check=True, capture_output=True)
v = openpyxl.load_workbook(tmp, data_only=True)
out = []
for sheet, col, rows in (("Gate plan", "M", plan), ("Gate events", "N", events), ("Conditions", "J", conds)):
    for r, (sid, _) in rows.items():
        out.append({"scenario": sid, "sheet": sheet, "row": r, "message": v[sheet][f"{col}{r}"].value})

# Register (AIG-INV-04) Evidence index, if it sits next to the Gate Log: the 12-month test for
# screening by reference (Playbook §4.6; scenarios F4b to F4d). Dates are relative to today.
reg_src = os.path.join(os.path.dirname(os.path.abspath(src)), "AIG-INV-04_AI_Register_Proposed.xlsx")
registers = []
if os.path.exists(reg_src):
    rtmp = os.path.join(os.path.dirname(tmp), "register-test.xlsx")
    shutil.copy(reg_src, rtmp)
    rwb = openpyxl.load_workbook(rtmp)
    rwb["AI Register"]["A4"] = "AIR-T001"
    today = datetime.datetime.combine(datetime.date.today(), datetime.time())
    days = lambda n: today + datetime.timedelta(days=n)
    ev = {
        4: ("F4b", dict(M=days(-500), N=days(200))),
        5: ("F4c", dict(M=days(-200), N=days(-10))),
        6: ("F4d", dict(M=days(-500), N=days(200), O="SH-0001 owner review")),
        7: ("F4d", dict(M=days(-60), N=days(300))),
    }
    E2 = rwb["Evidence index"]
    for r, (_, vals) in ev.items():
        put(E2, r, **{**dict(A=f"EV-T{r}", B="AIR-T001", C="DPIA", G="Accepted", H="link"), **vals})
    rwb.save(rtmp)
    subprocess.run([sys.executable, os.path.basename(recalc), rtmp], cwd=os.path.dirname(recalc), check=True, capture_output=True)
    rv = openpyxl.load_workbook(rtmp, data_only=True)["Evidence index"]
    for r, (sid, _) in ev.items():
        out.append({"scenario": sid, "sheet": "Evidence index", "row": r, "message": rv[f"L{r}"].value})
    registers.append(os.path.basename(reg_src) + " " + openpyxl.load_workbook(reg_src)["Document Control"]["B7"].value)

# Agent Record / ASBOM (AIG-AGT-04) Record QA, if it sits next to the Gate Log: the
# proportionate record levels (Playbook F.3; scenarios I7 to I10).
agt_src = os.path.join(os.path.dirname(os.path.abspath(src)), "AIG-AGT-04_Agent_Record_ASBOM_Proposed.xlsx")
if os.path.exists(agt_src):
    atmp = os.path.join(os.path.dirname(tmp), "asbom-test.xlsx")
    shutil.copy(agt_src, atmp)
    awb = openpyxl.load_workbook(atmp)
    AR = awb["Agent Record"]
    core = dict(A="AIR-T001", B="Test agent", C="Draft replies for officer review", H="Service Owner", I="Copilot Studio",
                M="T0 informational", S="Reads one mailbox; drafts replies", T="Service account SA-T1", U="Disable the flow",
                X="Delete drafts", AA="2027-10-01", AS="Pre-action approval", BB="1.0", BJ="A1 assisted", BK="DEC-T1",
                BA="DEC-UC-T1", AB="Approved", AC="Active", O="No", Q="No", P="None", R="No")
    ag = {
        6: ("I7", {}),
        7: ("I8", dict(U="")),
        8: ("I9", dict(T="Not disclosed by supplier")),
        9: ("I10", dict(M="T3 consequential agent")),
    }
    for r, (_, vals) in ag.items():
        put(AR, r, **{**core, **vals, "AY": f"AG-T{r}"})
    awb.save(atmp)
    subprocess.run([sys.executable, os.path.basename(recalc), atmp], cwd=os.path.dirname(recalc), check=True, capture_output=True)
    av = openpyxl.load_workbook(atmp, data_only=True)["Agent Record"]
    for r, (sid, _) in ag.items():
        out.append({"scenario": sid, "sheet": "Agent Record", "row": r, "message": av[f"BI{r}"].value})
    registers.append(os.path.basename(agt_src) + " " + openpyxl.load_workbook(agt_src)["Summary"]["B8"].value)
version = openpyxl.load_workbook(src)["Document Control"]["B7"].value
res = {"workbook": os.path.basename(src), "version": version, "also": registers, "date": datetime.date.today().isoformat(), "results": out}
here = os.path.join(os.path.dirname(__file__), "..", "results")
json.dump(res, open(os.path.join(here, "workbook-results.json"), "w"), indent=1, ensure_ascii=False)
open(os.path.join(here, "workbook-results.js"), "w").write("window.QA_WORKBOOK = " + json.dumps(res, ensure_ascii=False) + ";\n")
for o in out: print(o["scenario"], o["sheet"], o["row"], o["message"])
